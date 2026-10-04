"""Classificação de arquivos e extração de conteúdo.

Regras:
- O tipo é decidido primeiro pelo mimeType nativo do Google e depois pela extensão do
  nome (o Drive costuma marcar `.md` como text/plain ou application/octet-stream).
- Formato não suportado é "ignorado" com motivo; nunca se inventa conteúdo.
- Erro de leitura é "falha" (ReadError), nunca "arquivo vazio".
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .drive import FOLDER_MIME, GDOC_MIME, GSHEET_MIME, XLSX_MIME

MAX_BYTES = 10 * 1024 * 1024  # mesmo limite da exportação de Docs do Drive

SUPPORTED_KINDS = ("markdown", "text", "xlsx", "gdoc", "gsheet")

_IGNORED_BY_EXT = {
    ".docx": "Documento do Word não é lido diretamente. Para processar, abra no Drive com "
    "\"Documentos Google\" (o Google Doc gerado é lido).",
    ".doc": "Documento do Word não é lido diretamente. Converta para Google Docs.",
    ".pdf": "PDF ainda não processado nesta versão.",
    ".xls": "Planilha no formato antigo do Excel (.xls) não é lida; salve como .xlsx.",
    ".csv": "CSV ainda não processado nesta versão.",
    ".pptx": "Apresentação: formato ainda não processado.",
}
_IGNORED_BY_MIME_PREFIX = {
    "image/": "Imagem: formato ainda não processado (exigiria OCR). Nenhum conteúdo foi inferido.",
    "video/": "Vídeo: formato ainda não processado (exigiria transcrição). Nenhum conteúdo foi inferido.",
    "audio/": "Áudio: formato ainda não processado (exigiria transcrição). Nenhum conteúdo foi inferido.",
}
_GOOGLE_NATIVE_NAMES = {
    "application/vnd.google-apps.presentation": "Apresentação Google",
    "application/vnd.google-apps.form": "Formulário Google",
    "application/vnd.google-apps.drawing": "Desenho Google",
    "application/vnd.google-apps.shortcut": "Atalho do Drive (não seguido para não sair da pasta monitorada)",
}


class ReadError(Exception):
    """O arquivo existe mas não pôde ser lido. Vira sync_status = failed."""


@dataclass
class Extracted:
    text: str
    content_hash: str
    title: str | None = None
    meta: dict[str, str] = field(default_factory=dict)
    structured: dict[str, Any] | None = None
    note: str | None = None  # observação não fatal (ex.: arquivo sem conteúdo)


def file_ext(name: str) -> str:
    m = re.search(r"(\.[A-Za-z0-9]{1,8})$", name or "")
    return m.group(1).lower() if m else ""


def classify(name: str, mime_type: str) -> tuple[str | None, str | None]:
    """Devolve (kind, motivo_se_ignorado). kind None = não processado."""
    ext = file_ext(name)
    if mime_type == FOLDER_MIME:
        return "folder", None
    if mime_type == GDOC_MIME:
        return "gdoc", None
    if mime_type == GSHEET_MIME:
        return "gsheet", None
    if mime_type in _GOOGLE_NATIVE_NAMES:
        return None, f"{_GOOGLE_NATIVE_NAMES[mime_type]}: formato ainda não processado."
    if mime_type.startswith("application/vnd.google-apps."):
        return None, f"Formato nativo do Google ainda não processado ({mime_type.rsplit('.', 1)[-1]})."
    if ext in (".md", ".markdown") or mime_type == "text/markdown":
        return "markdown", None
    if ext == ".xlsx" or mime_type == XLSX_MIME:
        return "xlsx", None
    if ext == ".txt":
        return "text", None
    if ext in _IGNORED_BY_EXT:
        return None, _IGNORED_BY_EXT[ext]
    for prefix, reason in _IGNORED_BY_MIME_PREFIX.items():
        if mime_type.startswith(prefix):
            return None, reason
    label = ext or mime_type or "desconhecido"
    return None, f"Formato ainda não processado ({label})."


# ---------------------------------------------------------------------------
# Texto (Markdown, .txt, Google Docs exportado como text/plain)
# ---------------------------------------------------------------------------

_META_LINE = re.compile(r"^\s*([^\W\d]\w*)\s*:\s*(.+?)\s*$")


def _decode(data: bytes) -> str:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ReadError("O arquivo não está em UTF-8; não foi possível ler o texto.") from exc
    return text.replace("\r\n", "\n").replace("\r", "\n")


def parse_header(text: str) -> tuple[str | None, dict[str, str]]:
    """Título e o bloco de linhas `chave: valor` do início do documento.

    Aceita linhas em branco entre as linhas de cabeçalho (exportação de Google Docs) e
    texto antes do título: um Google Doc convertido de `.docx` exporta primeiro o cabeçalho
    de página (ex.: "LIGA IA UFSCAR / CASE TÉCNICO"). O bloco é procurado nas primeiras
    linhas; o título é o `# H1` ou a linha logo antes do bloco. Chaves sem espaço, como em
    `status: ativo` ou `data_da_reuniao: 2026-10-01`.
    """
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    if not lines:
        return None, {}
    h1_at = next((i for i, line in enumerate(lines[:3]) if line.startswith("# ")), None)
    # Com "# Título", o bloco vem logo abaixo dele; sem título marcado (exportação do Google Docs),
    # aceita até duas linhas antes (cabeçalho de página + título).
    candidates = [h1_at + 1] if h1_at is not None else [1, 2]
    start = next((i for i in candidates if i < len(lines) and _META_LINE.match(lines[i])), None)
    h1 = lines[h1_at].lstrip("#").strip() if h1_at is not None else None
    if start is None:
        title = h1 or lines[0].lstrip("#").strip() or None
        return title, {}
    meta: dict[str, str] = {}
    for line in lines[start:]:
        m = _META_LINE.match(line)
        if not m:
            break
        meta.setdefault(m.group(1).lower(), m.group(2))
    title = h1 or lines[start - 1].lstrip("#").strip() or None
    return title, meta


def _text_hash(text: str) -> str:
    # Espaços no fim da linha (Markdown usa "  " para quebra) não mudam o conteúdo.
    normalized = "\n".join(line.rstrip() for line in text.strip().split("\n"))
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def read_text(data: bytes) -> Extracted:
    text = _decode(data)
    title, meta = parse_header(text)
    note = "Arquivo sem conteúdo de texto." if not text.strip() else None
    return Extracted(text=text, content_hash=_text_hash(text), title=title, meta=meta, note=note)


# ---------------------------------------------------------------------------
# Planilhas (.xlsx e Google Sheets exportado como .xlsx)
# ---------------------------------------------------------------------------


def _cell(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        return value.strip() or None
    return value


def read_xlsx(data: bytes) -> Extracted:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, ValueError, OSError) as exc:
        raise ReadError(f"Planilha corrompida ou em formato inesperado ({exc.__class__.__name__}).") from exc

    sheets = []
    text_parts = []
    for ws in wb.worksheets:
        rows = []
        empty_rows = 0
        for idx, raw in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = [_cell(v) for v in raw]
            while cells and cells[-1] is None:
                cells.pop()
            if not cells:
                empty_rows += 1
                continue
            rows.append({"row": idx, "cells": cells})
        sheets.append({"name": ws.title, "rows": rows, "empty_rows": empty_rows})
        text_parts.append(f"## Aba: {ws.title}")
        for r in rows:
            values = " | ".join("" if c is None else str(c) for c in r["cells"])
            text_parts.append(f"Linha {r['row']}: {values}")
    wb.close()

    structured = {"sheets": [{"name": s["name"], "rows": s["rows"]} for s in sheets]}
    canonical = json.dumps(structured, ensure_ascii=False, sort_keys=True)
    data_rows = sum(max(len(s["rows"]) - 1, 0) for s in sheets)  # sem contar o cabeçalho
    note = None
    if data_rows == 0:
        note = "Planilha sem linhas de dados (só cabeçalho ou vazia)."
    structured["summary"] = {s["name"]: {"rows": len(s["rows"]), "empty_rows": s["empty_rows"]} for s in sheets}
    return Extracted(
        text="\n".join(text_parts),
        content_hash=hashlib.sha256(canonical.encode("utf-8")).hexdigest(),
        title=None,
        meta={},
        structured=structured,
        note=note,
    )


def extract(kind: str, data: bytes) -> Extracted:
    if len(data) > MAX_BYTES:
        raise ReadError("Arquivo maior que 10 MB; não processado.")
    if kind in ("markdown", "text", "gdoc"):
        return read_text(data)
    if kind in ("xlsx", "gsheet"):
        return read_xlsx(data)
    raise ValueError(f"tipo sem leitor: {kind}")
